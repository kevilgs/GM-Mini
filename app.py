"""Offline GM MINI workbook generator.

Run this file on the local PC, open http://127.0.0.1:5050, enter the
month's setup data, and download a self-contained GM MINI workbook.
"""
from __future__ import annotations

import calendar
import copy
import os
import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).with_name("vendor")))

from flask import Flask, abort, render_template, request, send_from_directory
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / "GM MINI- AUGUST 2026.xlsx"
OUTPUTS = ROOT / "generated_workbooks"

COMMODITIES = ["CEM", "COAL", "CONT", "FERT", "IMFT", "POL", "SALT", "STEEL", "DOC", "FG", "CHEM", "AUTO", "OTHERS"]
DIVISIONS = ["ADI", "GIMB", "BCT", "BRC", "RJT", "BVP", "RTM"]
RAKE_STOCKS = ["JUMBO", "BOXN", "BTPN", "BTPGN", "CONT", "SHRA", "OTHERS"]

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024


def column_for_day(first_column: int, day: int) -> str:
    return get_column_letter(first_column + day - 1)


def safe_number(raw: str | None, label: str) -> float:
    raw = (raw or "").strip().replace(",", "")
    if raw == "":
        return 0.0
    try:
        return float(raw)
    except ValueError as exc:
        raise ValueError(f"{label} must be a number.") from exc


def field_value(name: str, default: float = 0.0) -> float:
    return safe_number(request.form.get(name), name) if name in request.form else default


def template_defaults() -> dict:
    """Read useful starting values only; the source workbook is never changed."""
    wb = load_workbook(TEMPLATE, data_only=True, read_only=True)
    daily = wb["1"]
    link = wb["LINK"]
    defaults = {
        "commodity_targets": [daily.cell(5 + i, 2).value or 0 for i in range(len(COMMODITIES))],
        "division_targets": [daily.cell(21 + i, 2).value or 0 for i in range(len(DIVISIONS))],
        "current_baseline_commodities": [link.cell(3 + i, 101).value or 0 for i in range(len(COMMODITIES))],
        "current_baseline_divisions": [link.cell(20 + i, 101).value or 0 for i in range(len(DIVISIONS))],
        "prior_full_year_commodities": [link.cell(61 + i, 3).value or 0 for i in range(len(COMMODITIES))],
        "prior_full_year_divisions": [link.cell(77 + i, 3).value or 0 for i in range(len(DIVISIONS))],
        "annual_target_wagons": daily["Q5"].value or 0,
        "annual_target_mt": daily["S5"].value or 0,
        "prior_annual_wagons": daily["P5"].value or 0,
        "prior_annual_mt": daily["R5"].value or 0,
        "prior_progress_wagons": daily["P6"].value or 0,
        "prior_progress_mt": daily["R6"].value or 0,
    }
    wb.close()
    return defaults


def write_daily_formulas(ws, day: int) -> None:
    """Restore all formula links that vary by day number on a cloned daily sheet."""
    monthly_avg = column_for_day(68, day)       # BP
    historical_daily = column_for_day(2, day)   # B
    historical_mtd = column_for_day(35, day)    # AI
    current_fy_avg = column_for_day(137, day)   # EG
    historical_fy_avg = column_for_day(68, day) # BP
    current_fy_cumulative = column_for_day(104, day)  # CZ
    historical_month_cumulative = column_for_day(35, day)  # AI

    for index in range(len(COMMODITIES)):
        row = 5 + index
        link_row = 3 + index
        historical_row = 32 + index
        full_month_row = 61 + index
        ws.cell(row, 5).value = f"=LINK!{monthly_avg}{link_row}"
        ws.cell(row, 6).value = f"=LINK!{historical_daily}{historical_row}"
        ws.cell(row, 7).value = f"=LINK!{historical_mtd}{historical_row}"
        ws.cell(row, 8).value = f"=E{row}-G{row}"
        ws.cell(row, 9).value = f"=LINK!B{full_month_row}"
        ws.cell(row, 10).value = f"=E{row}-I{row}"
        ws.cell(row, 11).value = f"=LINK!{current_fy_avg}{link_row}"
        ws.cell(row, 12).value = f"=LINK!{historical_fy_avg}{historical_row}"
        ws.cell(row, 13).value = f"=LINK!C{full_month_row}"

    for index in range(len(DIVISIONS)):
        row = 21 + index
        link_row = 20 + index
        historical_row = 49 + index
        full_month_row = 77 + index
        ws.cell(row, 5).value = f"=LINK!{monthly_avg}{link_row}"
        ws.cell(row, 6).value = f"=LINK!{historical_daily}{historical_row}"
        ws.cell(row, 7).value = f"=LINK!{historical_mtd}{historical_row}"
        ws.cell(row, 8).value = f"=E{row}-G{row}"
        ws.cell(row, 9).value = f"=LINK!B{full_month_row}"
        ws.cell(row, 10).value = f"=E{row}-I{row}"
        ws.cell(row, 11).value = f"=LINK!{current_fy_avg}{link_row}"
        ws.cell(row, 12).value = f"=LINK!{historical_fy_avg}{historical_row}"
        ws.cell(row, 13).value = f"=LINK!C{full_month_row}"

    # Right-side summary panel. These cells retain the original layout.
    ws["P8"] = f"=LINK!{historical_month_cumulative}70"
    ws["Q8"] = f"=LINK!{monthly_avg}27"
    for idx, link_row in enumerate(range(20, 27), start=10):
        ws[f"Q{idx}"] = f"=LINK!{current_fy_cumulative}{link_row}"
    ws["Q16"] = f"=LINK!{current_fy_cumulative}27"


def clear_daily_inputs(ws) -> None:
    """Clear only operational input cells; labels, formatting and formulas remain."""
    for row in range(5, 18):
        ws.cell(row, 3).value = None
        ws.cell(row, 4).value = None
    for row in range(21, 28):
        ws.cell(row, 3).value = None
        ws.cell(row, 4).value = None
        ws.cell(row, 16).value = None

    # Daily operating detail below the dashboard is intentionally blank for a new month.
    for row in range(31, 83):
        for col in range(1, 23):
            cell = ws.cell(row, col)
            if cell.data_type != "f" and isinstance(cell.value, (int, float)):
                cell.value = None


def build_link(link, days: int, fiscal_days_before_month: int, values: dict) -> None:
    """Populate manual history and rebuild dynamic LINK formulas without altering its layout."""
    link["CX2"] = fiscal_days_before_month

    commodity_rows = list(range(3, 16))
    division_rows = list(range(20, 27))
    historic_commodity_rows = list(range(32, 45))
    historic_division_rows = list(range(49, 56))

    # Current-year opening balances through the end of the previous month.
    for row, amount in zip(commodity_rows, values["current_baseline_commodities"]):
        link.cell(row, 101).value = amount  # CW
    for row, amount in zip(division_rows, values["current_baseline_divisions"]):
        link.cell(row, 101).value = amount
    link["CW16"] = "=SUM(CW3:CW15)"
    link["CW27"] = "=SUM(CW20:CW26)"

    # Annual fiscal-day headers, e.g. 123 through 153 for August after 122 days.
    for day in range(1, 32):
        link.cell(2, 137 + day - 1).value = fiscal_days_before_month + day  # EG:FK

    # Current daily links and all current-year roll-forwards.
    for day in range(1, 32):
        daily_col = column_for_day(2, day)
        mtd_col = column_for_day(35, day)
        monthly_avg_col = column_for_day(68, day)
        fy_col = column_for_day(104, day)
        fy_avg_col = column_for_day(137, day)
        active = day <= days
        for source_row, daily_row in zip(commodity_rows, range(5, 18)):
            link[f"{daily_col}{source_row}"] = f"='{day}'!$D{daily_row}" if active else None
        for source_row, daily_row in zip(division_rows, range(21, 28)):
            link[f"{daily_col}{source_row}"] = f"='{day}'!$D{daily_row}" if active else None
        for source_row, daily_row in zip(range(89, 96), range(21, 28)):
            link[f"{daily_col}{source_row}"] = f"='{day}'!P{daily_row}" if active else None

        for row in commodity_rows + division_rows:
            if active:
                previous_mtd = column_for_day(35, day - 1) if day > 1 else None
                link[f"{mtd_col}{row}"] = f"={daily_col}{row}" if day == 1 else f"={previous_mtd}{row}+{daily_col}{row}"
                link[f"{monthly_avg_col}{row}"] = f"={mtd_col}{row}/{mtd_col}$2"
                previous_fy = column_for_day(104, day - 1) if day > 1 else "CW"
                link[f"{fy_col}{row}"] = f"={previous_fy}{row}+{daily_col}{row}"
                link[f"{fy_avg_col}{row}"] = f"={fy_col}{row}/{fy_avg_col}$2"
            else:
                for col in (mtd_col, monthly_avg_col, fy_col, fy_avg_col):
                    link[f"{col}{row}"] = None

        for total_row, start_row, end_row in ((16, 3, 15), (27, 20, 26), (96, 89, 95)):
            link[f"{daily_col}{total_row}"] = f"=SUM({daily_col}{start_row}:{daily_col}{end_row})" if active else None
            if active:
                link[f"{mtd_col}{total_row}"] = f"=SUM({mtd_col}{start_row}:{mtd_col}{end_row})"
                link[f"{monthly_avg_col}{total_row}"] = f"=SUM({monthly_avg_col}{start_row}:{monthly_avg_col}{end_row})"
                if total_row != 96:
                    link[f"{fy_col}{total_row}"] = f"=SUM({fy_col}{start_row}:{fy_col}{end_row})"
                    link[f"{fy_avg_col}{total_row}"] = f"=SUM({fy_avg_col}{start_row}:{fy_avg_col}{end_row})"

    # Manual prior-year daily grids. All prior-year MTD and FY averages are formulas.
    for day in range(1, 32):
        col = column_for_day(2, day)
        hist_mtd_col = column_for_day(35, day)
        hist_fy_avg_col = column_for_day(68, day)
        active = day <= days
        for index, row in enumerate(historic_commodity_rows):
            link[f"{col}{row}"] = values["prior_daily_commodities"][index][day - 1] if active else None
        for index, row in enumerate(historic_division_rows):
            link[f"{col}{row}"] = values["prior_daily_divisions"][index][day - 1] if active else None

        for row, baseline in zip(historic_commodity_rows, values["prior_baseline_commodities"]):
            if active:
                link[f"{hist_mtd_col}{row}"] = f"=SUM($B{row}:{col}{row})/{day}"
                link[f"{hist_fy_avg_col}{row}"] = f"=({baseline}+SUM($B{row}:{col}{row}))/({values['prior_fiscal_days']}+{day})"
            else:
                link[f"{hist_mtd_col}{row}"] = None
                link[f"{hist_fy_avg_col}{row}"] = None
        for row, baseline in zip(historic_division_rows, values["prior_baseline_divisions"]):
            if active:
                link[f"{hist_mtd_col}{row}"] = f"=SUM($B{row}:{col}{row})/{day}"
                link[f"{hist_fy_avg_col}{row}"] = f"=({baseline}+SUM($B{row}:{col}{row}))/({values['prior_fiscal_days']}+{day})"
            else:
                link[f"{hist_mtd_col}{row}"] = None
                link[f"{hist_fy_avg_col}{row}"] = None

        for total_row, start_row, end_row in ((45, 32, 44), (56, 49, 55)):
            link[f"{col}{total_row}"] = f"=SUM({col}{start_row}:{col}{end_row})" if active else None
            if active:
                link[f"{hist_mtd_col}{total_row}"] = f"=SUM({hist_mtd_col}{start_row}:{hist_mtd_col}{end_row})"
                link[f"{hist_fy_avg_col}{total_row}"] = f"=SUM({hist_fy_avg_col}{start_row}:{hist_fy_avg_col}{end_row})"

    last_col = column_for_day(2, days)
    for index in range(len(COMMODITIES)):
        link.cell(61 + index, 2).value = f"=AVERAGE(B{32 + index}:{last_col}{32 + index})"
        link.cell(61 + index, 3).value = values["prior_full_year_commodities"][index]
    link["B74"] = "=SUM(B61:B73)"
    link["C74"] = "=SUM(C61:C73)"
    for index in range(len(DIVISIONS)):
        link.cell(77 + index, 2).value = f"=AVERAGE(B{49 + index}:{last_col}{49 + index})"
        link.cell(77 + index, 3).value = values["prior_full_year_divisions"][index]
    link["B84"] = "=SUM(B77:B83)"
    link["C84"] = "=SUM(C77:C83)"

    # Right-side last-year summary panel: yearly cumulative divisions and month cumulative total.
    for day in range(1, 32):
        col = column_for_day(35, day)
        active = day <= days
        for index in range(len(DIVISIONS)):
            link_row = 60 + index
            source_row = 49 + index
            baseline = values["prior_baseline_divisions"][index]
            link[f"{col}{link_row}"] = f"={baseline}+SUM($B{source_row}:{column_for_day(2, day)}{source_row})" if active else None
        link[f"{col}67"] = f"=SUM({col}60:{col}66)" if active else None
        link[f"{col}70"] = f"=SUM($B45:{column_for_day(2, day)}45)" if active else None


def generate_workbook(year: int, month: int) -> Path:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template workbook is missing: {TEMPLATE.name}")
    days = calendar.monthrange(year, month)[1]
    first_day = date(year, month, 1)
    fiscal_start = date(year if month >= 4 else year - 1, 4, 1)
    fiscal_days_before_month = (first_day - fiscal_start).days

    values = {
        "commodity_targets": [field_value(f"commodity_target_{i}") for i in range(len(COMMODITIES))],
        "division_targets": [field_value(f"division_target_{i}") for i in range(len(DIVISIONS))],
        "current_baseline_commodities": [field_value(f"current_commodity_{i}") for i in range(len(COMMODITIES))],
        "current_baseline_divisions": [field_value(f"current_division_{i}") for i in range(len(DIVISIONS))],
        "prior_baseline_commodities": [field_value(f"prior_commodity_{i}") for i in range(len(COMMODITIES))],
        "prior_baseline_divisions": [field_value(f"prior_division_{i}") for i in range(len(DIVISIONS))],
        "prior_full_year_commodities": [field_value(f"prior_full_commodity_{i}") for i in range(len(COMMODITIES))],
        "prior_full_year_divisions": [field_value(f"prior_full_division_{i}") for i in range(len(DIVISIONS))],
        "prior_fiscal_days": field_value("prior_fiscal_days", fiscal_days_before_month),
        "prior_daily_commodities": [[field_value(f"prior_daily_commodity_{r}_{d}") for d in range(1, 32)] for r in range(len(COMMODITIES))],
        "prior_daily_divisions": [[field_value(f"prior_daily_division_{r}_{d}") for d in range(1, 32)] for r in range(len(DIVISIONS))],
    }

    wb = load_workbook(TEMPLATE, data_only=False)
    source = wb["1"]
    for sheet_name in [ws.title for ws in wb.worksheets if ws.title.isdigit() and ws.title != "1"]:
        del wb[sheet_name]
    for day in range(2, days + 1):
        ws = wb.copy_worksheet(source)
        ws.title = str(day)

    for day in range(1, days + 1):
        ws = wb[str(day)]
        report_day = date(year, month, day)
        ws["P1"] = report_day.fromordinal(report_day.toordinal() + 1)
        clear_daily_inputs(ws)
        for index, target in enumerate(values["commodity_targets"]):
            target_cell = ws.cell(5 + index, 2)
            if not isinstance(target_cell, MergedCell):
                target_cell.value = target
        for index, target in enumerate(values["division_targets"]):
            target_cell = ws.cell(21 + index, 2)
            if not isinstance(target_cell, MergedCell):
                target_cell.value = target
        write_daily_formulas(ws, day)

    link = wb["LINK"]
    build_link(link, days, fiscal_days_before_month, values)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    OUTPUTS.mkdir(exist_ok=True)
    output = OUTPUTS / f"GM MINI- {calendar.month_name[month].upper()} {year}.xlsx"
    wb.save(output)
    return output


@app.get("/")
def index():
    today = date.today()
    defaults = template_defaults()
    return render_template(
        "index.html",
        commodities=COMMODITIES,
        divisions=DIVISIONS,
        days=range(1, 32),
        months=list(enumerate(calendar.month_name))[1:],
        current_year=today.year,
        current_month=today.month,
        defaults=defaults,
    )


@app.post("/generate")
def generate():
    try:
        year = int(request.form["year"])
        month = int(request.form["month"])
        if not 2020 <= year <= 2100 or not 1 <= month <= 12:
            raise ValueError("Choose a valid month and year.")
        output = generate_workbook(year, month)
        return render_template("complete.html", filename=output.name, days=calendar.monthrange(year, month)[1])
    except (ValueError, FileNotFoundError) as exc:
        return render_template("error.html", message=str(exc)), 400


@app.get("/download/<path:filename>")
def download(filename: str):
    if Path(filename).name != filename:
        abort(404)
    return send_from_directory(OUTPUTS, filename, as_attachment=True)


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5050, debug=False)
